"""Tests for z7_officeletters.core.documents."""

from __future__ import annotations

import pytest

from z7_officeletters.core.documents import (
    _titlecase_nome,
    construir_nome_arquivo,
    normalizar_numero_mocao,
    criar_modelo_envelope,
)


# =============================================================================
# normalizar_numero_mocao
# =============================================================================
class TestNormalizarNumeroMocao:

    def test_numero_puro_nao_alterado(self) -> None:
        assert normalizar_numero_mocao("124") == "124"

    def test_remove_espacos(self) -> None:
        assert normalizar_numero_mocao("  124  ") == "124"

    def test_sufixo_nao_numerico_preservado(self) -> None:
        assert normalizar_numero_mocao("124-A") == "124-A"

    @pytest.mark.parametrize("entrada,esperado", [
        ("124/2026", "124"),
        ("124-2026", "124"),
        ("124/26",   "124"),
        ("124-26",   "124"),
        ("001/2026", "001"),
        ("999-2025", "999"),
        ("7",        "7"),
    ])
    def test_variantes(self, entrada: str, esperado: str) -> None:
        assert normalizar_numero_mocao(entrada) == esperado


# =============================================================================
# construir_nome_arquivo
# =============================================================================
class TestConstruirNomeArquivo:

    def _nome(self, **overrides: object) -> str:
        params: dict[str, object] = dict(
            num_oficio_str="001",
            sigla_servidor="js",
            tipo_mocao="Aplauso",
            num_mocao="124",
            envio="E-mail",
            nome_dest="Fulano de Tal",
            sigla_autores="AD",
            ano=2026,
        )
        params.update(overrides)
        return construir_nome_arquivo(**params)  # type: ignore[arg-type]

    def test_extensao_docx(self) -> None:
        assert self._nome().endswith(".docx")

    def test_contem_numero_oficio(self) -> None:
        assert "001" in self._nome()

    def test_contem_tipo_mocao(self) -> None:
        assert "Aplauso" in self._nome()

    def test_contem_numero_mocao_com_sufixo_26(self) -> None:
        assert "124-26" in self._nome()

    def test_sufixo_26_aparece_uma_vez(self) -> None:
        assert self._nome().count("-26") == 1

    def test_envio_convertido_para_minusculo(self) -> None:
        assert "e-mail" in self._nome(envio="E-mail")

    def test_sigla_servidor_refletida(self) -> None:
        assert "redator" in self._nome(sigla_servidor="redator")

    def test_sigla_autores_refletida(self) -> None:
        assert "ad e outros" in self._nome(sigla_autores="ad e outros")

    def test_remove_caracteres_invalidos_windows(self) -> None:
        nome = construir_nome_arquivo(
            num_oficio_str="001",
            sigla_servidor="js",
            tipo_mocao="Aplauso",
            num_mocao="124",
            envio="Em Mãos",
            nome_dest='Nome "Ilegal" <teste>',
            sigla_autores="AD",
            ano=2026,
        )
        for ch in r'\/*?:"<>|':
            assert ch not in nome

    def test_nome_dest_caps_convertido_para_titulo(self) -> None:
        assert "Marcos Silva" in self._nome(nome_dest="MARCOS SILVA")

    def test_nome_dest_preposicao_minuscula(self) -> None:
        assert "João da Silva" in self._nome(nome_dest="JOÃO DA SILVA")

    def test_nome_dest_instituicao_caps_titulo(self) -> None:
        assert "Câmara Municipal" in self._nome(nome_dest="CÂMARA MUNICIPAL")

    def test_nome_dest_abreviacao_com_ponto_preservada(self) -> None:
        assert "S.A." in self._nome(nome_dest="EMPRESA S.A.")

    def test_nome_dest_ja_em_titulo_nao_alterado(self) -> None:
        assert "Fulano de Tal" in self._nome(nome_dest="Fulano de Tal")


# =============================================================================
# _titlecase_nome
# =============================================================================
class TestTitlecaseNome:

    def test_nome_pessoa_caps(self) -> None:
        assert _titlecase_nome("MARCOS SILVA") == "Marcos Silva"

    def test_preposicoes_minusculas(self) -> None:
        assert _titlecase_nome("JOÃO DA SILVA") == "João da Silva"

    def test_multiplas_preposicoes(self) -> None:
        assert _titlecase_nome("ASSOCIAÇÃO DOS MORADORES DE BAIRRO") == "Associação dos Moradores de Bairro"

    def test_preposicao_primeira_palavra_preservada(self) -> None:
        # First token is never lowercased even if it is a preposition
        assert _titlecase_nome("DE SOUZA").startswith("De")

    def test_abreviacao_com_ponto_preservada(self) -> None:
        assert _titlecase_nome("EMPRESA S.A.") == "Empresa S.A."

    def test_acronimo_sem_vogal_preservado(self) -> None:
        # BNB = B-N-B, no vowels → kept uppercase
        assert _titlecase_nome("AG\u00caNCIA BNB") == "Ag\u00eancia BNB"
        # CNPJ = C-N-P-J, no vowels → kept uppercase
        assert _titlecase_nome("CNPJ") == "CNPJ"

    def test_string_vazia(self) -> None:
        assert _titlecase_nome("") == ""

    def test_uma_palavra(self) -> None:
        assert _titlecase_nome("PREFEITURA") == "Prefeitura"

    def test_acronimo_com_vogal_preservado(self) -> None:
        # APAE, OAB, SUS contain vowels but should be preserved in uppercase
        assert _titlecase_nome("ASSOCIAÇÃO OAB") == "Associação OAB"
        assert _titlecase_nome("APAE DE SANTA BÁRBARA") == "APAE de Santa Bárbara"
        assert _titlecase_nome("SISTEMA SUS") == "Sistema SUS"


# =============================================================================
# criar_modelo_envelope
# =============================================================================
class TestCriarModeloEnvelope:

    def test_cria_envelope_com_dimensoes_e_placeholders(self, tmp_path: Path) -> None:
        from docx import Document  # type: ignore[import-untyped]
        
        caminho_temp = tmp_path / "modelo_envelope.docx"
        assert not caminho_temp.exists()

        resultado = criar_modelo_envelope(caminho_temp)
        assert resultado == caminho_temp
        assert caminho_temp.exists()

        # Load it and verify size and sections
        doc = Document(caminho_temp)
        section = doc.sections[0]
        
        # Verify DL dimensions: 22 cm by 11 cm
        assert abs(section.page_width.cm - 22.0) < 0.01
        assert abs(section.page_height.cm - 11.0) < 0.01

        # Read text and verify placeholders exist in the text
        texto_completo = "\n".join(p.text for p in doc.paragraphs)
        assert "TRATAMENTO_RODAPE" in texto_completo
        assert "DESTINATARIO_NOME" in texto_completo
        assert "DESTINATARIO_ENDERECO" in texto_completo
        assert "REMETENTE" in texto_completo
